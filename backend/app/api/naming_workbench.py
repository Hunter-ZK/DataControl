"""Authenticated, bounded Excel import/export and deterministic naming preview."""
from __future__ import annotations

import csv
import io
from collections import Counter
from zipfile import BadZipFile, ZipFile
from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from backend.app.api.assets import get_db
from backend.app.api.auth import get_current_user
from backend.app.db.models import User
from backend.app.services.naming import suggest_names

router = APIRouter(prefix="/api/v1/workbench/naming", tags=["naming-workbench"])
MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 80

class NamingRequest(BaseModel):
    names: list[str] = Field(min_length=1, max_length=MAX_ROWS)
    maxLength: int = Field(default=64, ge=8, le=128)

class ExportEntry(BaseModel):
    original: dict[str, str] = Field(default_factory=dict)
    cnName: str = Field(max_length=200)
    suggestedName: str = Field(default="", max_length=128)
    segments: str = Field(default="", max_length=2000)
    source: str = Field(default="", max_length=100)
    issues: str = Field(default="", max_length=2000)
    status: str = Field(max_length=50)
    note: str = Field(default="", max_length=2000)

class ExportRequest(BaseModel):
    entries: list[ExportEntry] = Field(min_length=1, max_length=MAX_ROWS)
    originalHeaders: list[str] = Field(default_factory=list, max_length=MAX_COLUMNS)
    rootSuggestions: list[dict[str, str | int]] = Field(default_factory=list, max_length=MAX_ROWS)

def _safe_excel(value):
    text = str(value if value is not None else "")
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + text
    return text[:32700]

def _read_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig")
    return [[str(v or "") for v in row] for row in csv.reader(io.StringIO(text))]

@router.post("/import")
async def import_file(file: UploadFile, _: User = Depends(get_current_user)):
    filename = (file.filename or "").lower()
    if not (filename.endswith(".xlsx") or filename.endswith(".csv")):
        raise HTTPException(422, "Only .xlsx and .csv files are supported")
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File exceeds 8 MiB limit")
    try:
        if filename.endswith(".xlsx"):
            with ZipFile(io.BytesIO(data)) as zip_file:
                records = zip_file.infolist()
                if len(records) > 3000 or sum(x.file_size for x in records) > 50 * 1024 * 1024:
                    raise HTTPException(413, "Workbook decompression size exceeds safety limit")
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheet = wb.active
                rows = [[str(v if v is not None else "") for v in cells[:MAX_COLUMNS+1]]
                        for cells in sheet.iter_rows(min_row=1, max_row=MAX_ROWS+2, values_only=True)]
            finally:
                wb.close()
        else:
            rows = _read_csv(data)
    except (ValueError, UnicodeDecodeError, BadZipFile, OSError, KeyError) as exc:
        raise HTTPException(422, "Invalid or unreadable spreadsheet") from exc
    if not rows:
        raise HTTPException(422, "Spreadsheet is empty")
    if len(rows) > MAX_ROWS + 1:
        raise HTTPException(422, "Row limit exceeded (2000)")
    if any(len(row) > MAX_COLUMNS for row in rows):
        raise HTTPException(422, "Column limit exceeded (80)")
    headers = [str(x).strip() for x in rows[0]]
    if not any(headers) or len(set(headers)) != len(headers) or any(not h for h in headers):
        raise HTTPException(422, "Column headers must be nonempty and unique")
    entries = [dict(zip(headers, row)) for row in rows[1:] if any(v.strip() for v in row)]
    return {"code": "OK", "data": {"headers": headers, "items": entries, "count": len(entries)}}

@router.post("/suggest")
def suggest(body: NamingRequest, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return {"code": "OK", "data": {
        "items": suggest_names(db, body.names, max_length=body.maxLength),
        "engine": "deterministic-root-matcher-v1",
        "notice": "规则预览：无模型臆测；需以正式命名规范与已确认词根审核后采纳。",
    }}

@router.post("/export")
def export_file(body: ExportRequest, _: User = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "命名结果"
    extra = [name for name in body.originalHeaders if name not in {"中文名", "建议英文名"}]
    ws.append(["序号", "中文名", "建议英文名", "拆分过程", "来源", "校验结果", "状态", "备注", *extra])
    for index, item in enumerate(body.entries, 1):
        ws.append([index, _safe_excel(item.cnName), _safe_excel(item.suggestedName),
                   _safe_excel(item.segments), _safe_excel(item.source), _safe_excel(item.issues),
                   _safe_excel(item.status), _safe_excel(item.note),
                   *[_safe_excel(item.original.get(name)) for name in extra]])
    roots = wb.create_sheet("建议新增词根")
    roots.append(["中文词", "建议英文缩写", "英文全称", "出现次数", "示例字段"])
    for suggestion in body.rootSuggestions:
        roots.append([_safe_excel(suggestion.get("cn", "")), _safe_excel(suggestion.get("abbr", "")),
                      _safe_excel(suggestion.get("full", "")), suggestion.get("count", 0),
                      _safe_excel(suggestion.get("examples", ""))])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            from openpyxl.styles import Font, PatternFill
            cell.fill = PatternFill("solid", fgColor="E9F2FF")
            cell.font = Font(bold=True, color="204476")
        for col in sheet.columns:
            from openpyxl.utils import get_column_letter
            letter = get_column_letter(col[0].column)
            sheet.column_dimensions[letter].width = min(45, max(13, max(len(str(c.value or "")) for c in col[:100]) + 3))
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="datacontrol_naming_results.xlsx"',
                 "Cache-Control": "no-store"})
