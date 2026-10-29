from io import BytesIO
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app.services.naming import suggest_names
from backend.app.db.session import SessionLocal

client = TestClient(app)

def _login():
    response = client.post("/api/v1/auth/login", json={"username": "demo", "password": "DataControl123!"})
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["data"]["accessToken"]}

def test_naming_requires_login():
    assert client.post("/api/v1/workbench/naming/suggest", json={"names": ["贷款余额"]}).status_code in (401, 403)

def test_deterministic_suggestions():
    with SessionLocal() as db:
        first = suggest_names(db, ["贷款余额", "这个词根绝对不存在"])
        second = suggest_names(db, ["贷款余额", "这个词根绝对不存在"])
        assert first == second
        assert second[1]["status"] == "待确认"
        assert second[1]["issues"]

def test_excel_roundtrip_and_escape():
    headers = _login()
    wb = Workbook()
    ws = wb.active
    ws.append(["中文名", "所属表"])
    ws.append(["贷款余额", "dws_loan"])
    data = BytesIO()
    wb.save(data)
    response = client.post("/api/v1/workbench/naming/import", headers=headers,
        files={"file": ("fields.xlsx", data.getvalue(),
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    assert response.json()["data"]["items"][0]["中文名"] == "贷款余额"
    export = client.post("/api/v1/workbench/naming/export", headers=headers, json={
        "originalHeaders": ["中文名", "所属表"],
        "entries": [{"original": {"所属表": "dws_loan"}, "cnName": "贷款余额",
                     "suggestedName": "=HYPERLINK(1)", "status": "待确认"}],
        "rootSuggestions": [{"cn": "某项", "abbr": "", "full": "", "count": 1, "examples": "贷款余额"}],
    })
    assert export.status_code == 200
    result = load_workbook(BytesIO(export.content), read_only=True)
    assert result.sheetnames == ["命名结果", "建议新增词根"]
    assert result.active["C2"].value.startswith("'=")
    assert result.active["I2"].value == "dws_loan"
